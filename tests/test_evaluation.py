import io
import uuid

from openpyxl import load_workbook

from tests.conftest import STORAGE_DIR, make_pdf, make_trf_pdf

API = "/api/v1"
TRF_ROWS = [
    ("5", "GENERAL CONDITIONS FOR THE TESTS", "P"),
    ("6", "CLASSIFICATION", "P"),
    ("6.1", "Protection against electric shock", "P"),
    ("7", "MARKING AND INSTRUCTIONS", "P"),
    ("7.1", "Rated voltage or voltage range", "P"),
    ("10", "POWER INPUT AND CURRENT", "P"),
    ("10.1", "Power input at normal operating temperature", "P"),
    ("15.3", "Humidity treatment", "N/A"),
]


def _std(client, text):
    pdf = make_pdf(text + f"\nuid {uuid.uuid4()}")
    return client.post(f"{API}/standards", files={"file": ("s.pdf", pdf, "application/pdf")}).json()["standard"]


def _report():
    return ("wm report.pdf", make_trf_pdf("Standard: IEC 60335-2-7:2012  uid " + str(uuid.uuid4()), TRF_ROWS),
            "application/pdf")


def test_evaluation_end_to_end(client, fake_ai):
    p1 = _std(client, "IEC 60335-1:2010")
    p2 = _std(client, "IEC 60335-2-7:2012 use with IEC 60335-1:2010")

    r = client.post(f"{API}/evaluations", data={"user_id": "amit"}, files={"report_file": _report()})
    assert r.status_code == 200, r.text
    ev = r.json()
    assert ev["from_cache"] is False
    assert ev["standards_used"] == ["IEC 60335-1:2010", "IEC 60335-2-7:2012"]
    assert ev["final_verdict"].startswith("REPORT CORRECTION")
    assert ev["counts"]["clause_rows"] == len(TRF_ROWS)
    assert len(fake_ai.evals) == 1 and len(fake_ai.evals[0]["files"]) == 3
    assert "10.1 | P |" in fake_ai.evals[0]["text"]                      # clause register sent to AI

    # files on disk next to the report
    pdf_path = STORAGE_DIR / ev["pdf_path"]
    xlsx_path = STORAGE_DIR / ev["xlsx_path"]
    assert pdf_path.exists() and pdf_path.read_bytes().startswith(b"%PDF")
    assert f"reports/{ev['report_id']}/outputs/evaluation_{ev['evaluation_id']}/" in ev["pdf_path"]

    wb = load_workbook(xlsx_path)
    assert wb.sheetnames == ["Executive Summary", "Critical Findings", "Part 2-7 Coverage",
                             "Part 1 Corrections", "Clause-by-Clause Review"]
    reg = {row[0].value: row for row in wb["Clause-by-Clause Review"].iter_rows(min_row=2)}
    assert reg["10.1"][3].value == "CRITICAL / SPECIFIC REVIEW"
    assert reg["10.1"][4].value == "CRITICAL: 400 W vs 2000 W cover"
    assert reg["10"][3].value == "REVIEW—PARENT"
    assert reg["6.1"][3].value == "REPORTED PASS—NOT VERIFIED TO WM-1"
    assert reg["15.3"][3].value == "REPORTED N/A—APPLICABILITY TO CONFIRM"
    assert reg["7.1"][4].value == "Marking photo missing"

    # downloads
    d = client.get(ev["download_pdf"])
    assert d.status_code == 200 and d.headers["content-type"] == "application/pdf"
    d = client.get(ev["download_xlsx"])
    assert d.status_code == 200 and load_workbook(io.BytesIO(d.content))

    # same report + standards -> cached, AI not called again
    r2 = client.post(f"{API}/evaluations", data={"user_id": "amit", "report_id": ev["report_id"]}).json()
    assert r2["from_cache"] is True and r2["evaluation_id"] == ev["evaluation_id"]
    assert len(fake_ai.evals) == 1

    # re-render without AI
    rr = client.post(f"{API}/evaluations/{ev['evaluation_id']}/render", data={"user_id": "amit"})
    assert rr.status_code == 200 and len(fake_ai.evals) == 1

    assert client.get(f"{API}/evaluations", params={"user_id": "amit"}).json()[0]["evaluation_id"] == ev["evaluation_id"]
    assert client.get(f"{API}/evaluations/{ev['evaluation_id']}", params={"user_id": "other"}).status_code == 403


def test_evaluation_from_chat(client, fake_ai):
    _std(client, "IEC 60335-1:2010")
    _std(client, "IEC 60335-2-7:2012")
    chat = client.post(f"{API}/chat", data={"user_id": "u", "query": "hi"}, files={"report_file": _report()}).json()
    r = client.post(f"{API}/evaluations", data={"user_id": "u", "chat_id": chat["chat_id"]})
    assert r.status_code == 200, r.text
    assert len(r.json()["standards_used"]) == 2


def test_evaluation_without_standards_is_rejected(client):
    r = client.post(f"{API}/evaluations", data={"user_id": "u"}, files={"report_file": _report()})
    assert r.status_code == 400
    assert "IEC 60335-2-7" in r.json()["error"]["message"]


def test_evaluation_needs_a_report(client):
    r = client.post(f"{API}/evaluations", data={"user_id": "u"})
    assert r.status_code == 400
