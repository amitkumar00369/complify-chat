"""Evaluation workbook in the same 5-sheet layout as the client's golden evaluation."""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.config.settings import get_settings
from app.schema.evaluation import EvaluationDocument

THIN = Side(style="thin", color="FFC9D3DE")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")

VERDICT_FILL = {
    "HOLD": "FFF8D7DA", "FAIL": "FFF8D7DA", "NON-CONFORMITY": "FFF8D7DA", "CORRECT": "FFFCE5CD",
    "CONDITIONAL": "FFFFF2CC", "INSUFFICIENT": "FFFFF2CC", "REVIEW": "FFFFF2CC",
    "PASS": "FFD9EAD3", "ACCEPTED": "FFD9EAD3",
}
RISK_FILL = {"Critical": "FFF8D7DA", "High": "FFFCE5CD", "Medium": "FFFFF2CC", "Low": "FFD9EAD3",
             **{f"P0-0{i}": "FFF8D7DA" for i in range(1, 6)}, "P1": "FFFCE5CD", "P2": "FFFFF2CC"}


def _hex(color: str) -> str:
    return "FF" + color.lstrip("#").upper()


def _verdict_fill(text: str) -> PatternFill | None:
    t = (text or "").upper()
    for key, color in VERDICT_FILL.items():   # order matters: HOLD before PASS
        if key in t:
            return PatternFill("solid", fgColor=color)
    return None


def _table(ws, headers: list[str], rows: list[list], widths: list[int], header_fill: PatternFill,
           color_col: int | None = None, risk_col: int | None = None) -> None:
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True, color="FFFFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = BORDER
    for r in rows:
        ws.append(r)
        i = ws.max_row
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=i, column=c)
            cell.alignment = WRAP
            cell.border = BORDER
        if color_col is not None:
            fill = _verdict_fill(str(r[color_col]))
            if fill:
                ws.cell(row=i, column=color_col + 1).fill = fill
        if risk_col is not None and r[risk_col] in RISK_FILL:
            ws.cell(row=i, column=risk_col + 1).fill = PatternFill("solid", fgColor=RISK_FILL[r[risk_col]])
    for idx, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"


def render_xlsx(doc: EvaluationDocument) -> bytes:
    s = get_settings()
    brand = _hex(s.brand_color)
    hfill = PatternFill("solid", fgColor=brand)
    ai = doc.ai
    wb = Workbook()

    # ---- Executive Summary ----
    ws = wb.active
    ws.title = "Executive Summary"
    ws.merge_cells("A1:H1")
    ws["A1"] = f"EVALUATION CONCLUSIONS FINDINGS — {ai.title_product.upper()}"
    ws["A1"].font = Font(size=15, bold=True, color="FFFFFFFF")
    ws["A1"].fill = hfill
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 26
    ws.append([])
    label_font = Font(bold=True, color=brand)
    kv = [
        ("Report number", f"{ai.report_number} | {ai.laboratory} | {ai.report_date}"),
        ("Product declared", ai.product_declared),
        ("Standards claimed", ai.standards_claimed),
        ("Evaluation", ai.evaluation_basis),
        ("Final verdict", ai.final_verdict),
        ("Reason", ai.reason),
    ]
    for k, v in kv:
        ws.append([k, v])
        r = ws.max_row
        ws.cell(row=r, column=1).font = label_font
        ws.cell(row=r, column=2).alignment = WRAP
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        if k == "Final verdict":
            ws.cell(row=r, column=2).font = Font(bold=True)
            fill = _verdict_fill(v)
            if fill:
                ws.cell(row=r, column=2).fill = fill
    if ai.product_facts:
        ws.append([])
        ws.append(["Product Facts (basis for applicability)"])
        r = ws.max_row
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
        ws.cell(row=r, column=1).font = Font(bold=True, color="FFFFFFFF")
        ws.cell(row=r, column=1).fill = hfill
        for pf in ai.product_facts:
            ws.append([pf.fact, pf.value, pf.source])
            r = ws.max_row
            ws.cell(row=r, column=1).font = label_font
            for c in (2, 3):
                ws.cell(row=r, column=c).alignment = WRAP
    ws.append([])
    ws.append(["Major Non-Conformity / Report-Validity Decision Points"])
    r = ws.max_row
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
    ws.cell(row=r, column=1).font = Font(bold=True, color="FFFFFFFF")
    ws.cell(row=r, column=1).fill = hfill
    for dp in ai.decision_points:
        ws.append([f"{dp.priority} — {dp.title}", dp.finding, dp.required_resolution])
        r = ws.max_row
        ws.cell(row=r, column=1).font = label_font
        for c in (1, 2, 3):
            ws.cell(row=r, column=c).alignment = WRAP
    ws.append([])
    ws.append(["Final Technical Conclusion"])
    r = ws.max_row
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
    ws.cell(row=r, column=1).font = Font(bold=True, color="FFFFFFFF")
    ws.cell(row=r, column=1).fill = hfill
    ws.append([ai.final_technical_conclusion])
    r = ws.max_row
    ws.merge_cells(start_row=r, start_column=1, end_row=r + 6, end_column=8)
    ws.cell(row=r, column=1).alignment = WRAP
    for col, w in zip("ABCDE", (31, 82, 72, 65, 14)):
        ws.column_dimensions[col].width = w

    # ---- Critical Findings ----
    _table(wb.create_sheet("Critical Findings"),
           ["Root cause", "Clause / concern", "Report condition", "Evaluation Engineer comments", "Required verdict",
            "Required action", "Report reference"],
           [[f.root_cause, f.clause_or_concern, f.report_condition, f.engineer_comments, f.required_verdict,
             f.required_action, f.report_reference] for f in ai.critical_findings],
           [11, 31, 65, 72, 28, 70, 22], hfill, color_col=4, risk_col=0)

    # ---- Part 2-x Coverage ----
    _table(wb.create_sheet(f"{doc.part2_label} Coverage"[:31]),
           ["Clause(s)", "Safety area", "Report status", "Applicability rule", "Required report outcome",
            "Risk level"],
           [[c.clauses, c.safety_area, c.report_status, c.applicability_rule, c.required_outcome,
             c.risk_level] for c in ai.part2_coverage],
           [24, 37, 45, 46, 55, 16], hfill, risk_col=5)

    # ---- Part 1 Corrections ----
    _table(wb.create_sheet("Part 1 Corrections"),
           ["Clause / issue", "Evaluation finding", "Evaluation Engineer comments", "Required correction",
            "Final required status"],
           [[c.clause_or_issue, c.finding, c.engineer_comments, c.required_correction, c.final_status]
            for c in ai.part1_corrections],
           [35, 58, 65, 58, 26], hfill, color_col=4)

    # ---- Clause-by-Clause Review ----
    _table(wb.create_sheet("Clause-by-Clause Review"),
           ["Clause", "Report row / topic", "Report verdict", "Preliminary assessment",
            "Required verification / discrepancy", "PDF page", "Evidence status"],
           [[r.clause, r.topic, r.report_verdict, r.assessment, r.required_verification, r.pdf_page,
             r.evidence_status] for r in doc.clause_review],
           [15, 71, 19, 41, 62, 15, 46], hfill)
    crit = wb["Clause-by-Clause Review"]
    for row in crit.iter_rows(min_row=2):
        if row[3].value and row[3].value.startswith("CRITICAL"):
            for cell in row:
                cell.fill = PatternFill("solid", fgColor="FFFFF2CC")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
